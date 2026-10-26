"""Pruebas de integración: la API real, los repositorios SQL y las migraciones de Alembic contra
PostgreSQL, sin repositorios falsos. Usan la base espodi_test (o TEST_DATABASE_URL), nunca la de trabajo.
Solo OpenAI se reemplaza, porque es un servicio externo que cobra por uso y no responde siempre igual.

Se corren aparte de las unitarias:  pytest -m integracion
"""
import os
from io import BytesIO
from pathlib import Path

import pytest
from alembic import command
from alembic.config import Config
from fastapi.testclient import TestClient
from openpyxl import Workbook
from sqlalchemy import create_engine, text
from sqlalchemy.engine import URL, make_url
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import sessionmaker

from src.compartido.bd import Base, get_db
from src.compartido.configuracion import settings
from src.main import app
from src.modulos.acceso_roles.aplicacion.asegurar_administrador import asegurar_administrador
from src.modulos.acceso_roles.infraestructura.repositorio_usuarios import RepositorioUsuariosSQL

CARPETA_BACKEND = Path(__file__).resolve().parents[2]
PASSWORD_INICIAL = settings.default_technician_password
ADMIN_INICIAL = ("admin", "Admin#Inicial1")  # como ADMIN_PASSWORD_INICIAL en .env
ADMIN = ("admin", "Admin#2026")


def url_de_prueba() -> URL:
    url = make_url(os.environ.get("TEST_DATABASE_URL") or make_url(settings.database_url).set(database="espodi_test"))
    # protección: estas pruebas vacían la base, así que nunca se usa una que no sea de pruebas
    if not (url.database or "").endswith("_test"):
        pytest.exit(f"La base de las pruebas de integración debe terminar en _test, no '{url.database}'")
    return url


def configuracion_alembic(url: URL) -> Config:
    config = Config(str(CARPETA_BACKEND / "alembic.ini"))
    config.attributes["url"] = url.render_as_string(hide_password=False)
    return config


@pytest.fixture(scope="session")
def motor():
    url = url_de_prueba()
    motor = create_engine(url)
    try:
        motor.connect().close()
    except OperationalError:
        pytest.exit(f"No se pudo conectar a la base '{url.database}'. ¿Está PostgreSQL encendido y la base creada?")

    # se parte de una base vacía y se aplican todas las migraciones, igual que en una instalación nueva
    config = configuracion_alembic(url)
    command.downgrade(config, "base")
    command.upgrade(config, "head")
    yield motor
    motor.dispose()


@pytest.fixture(autouse=True)
def base_vacia(motor):
    """Cada prueba empieza sin datos, para que no dependa del orden en que se corren"""
    tablas = ", ".join(t.name for t in Base.metadata.sorted_tables)
    with motor.begin() as conexion:
        conexion.execute(text(f"TRUNCATE {tablas} RESTART IDENTITY CASCADE"))


@pytest.fixture
def sesiones(motor):
    return sessionmaker(bind=motor, autoflush=False)


@pytest.fixture
def api(sesiones, tmp_path, monkeypatch):
    def get_db_de_prueba():
        sesion = sesiones()
        try:
            yield sesion
        finally:
            sesion.close()

    # las imágenes de los comunicados van a una carpeta temporal, no a datos/
    monkeypatch.setattr(settings, "comunicados_imagenes_dir", str(tmp_path / "imagenes"))
    app.dependency_overrides[get_db] = get_db_de_prueba
    # como al arrancar el backend: se crea el administrador inicial, que hace su primer acceso
    with sesiones() as sesion:
        asegurar_administrador(*ADMIN_INICIAL, RepositorioUsuariosSQL(sesion))
    api = Api(TestClient(app))
    api.primer_acceso(ADMIN[0], ADMIN[1], password_actual=ADMIN_INICIAL[1])
    yield api
    app.dependency_overrides.clear()


class Api:
    """Cliente HTTP con los pasos que se repiten en los flujos: iniciar sesión, primer acceso, etc."""

    def __init__(self, cliente: TestClient):
        self.cliente = cliente

    def login(self, usuario: str, password: str) -> dict:
        respuesta = self.cliente.post("/api/v1/auth/login", json={"nombre_usuario": usuario, "password": password})
        assert respuesta.status_code == 200, respuesta.text
        return {"Authorization": f"Bearer {respuesta.json()['access_token']}"}

    def admin(self) -> dict:
        return self.login(*ADMIN)

    def primer_acceso(self, usuario: str, password_nueva: str, password_actual: str = PASSWORD_INICIAL) -> dict:
        """Entra con la contraseña inicial, elige una propia y vuelve a entrar con ella"""
        cabecera = self.login(usuario, password_actual)
        respuesta = self.cliente.patch(
            "/api/v1/perfil/primer-acceso", headers=cabecera,
            json={"nombre_usuario": usuario, "password_nueva": password_nueva},
        )
        assert respuesta.status_code == 200, respuesta.text
        return self.login(usuario, password_nueva)

    def crear_supervisor(self, nombre: str, apellido: str, grupo: str) -> tuple[dict, dict]:
        """Lo registra el administrador; devuelve el supervisor y su cabecera de sesión"""
        respuesta = self.cliente.post("/api/v1/supervisores", headers=self.admin(), json={
            "nombre": nombre, "apellido_paterno": apellido, "nombre_de_grupo": grupo,
            "horario_entrada": "07:00", "horario_salida": "15:00",
        })
        assert respuesta.status_code == 201, respuesta.text
        supervisor = respuesta.json()
        return supervisor, self.primer_acceso(supervisor["nombre_usuario"], "Superv1sor!")

    def crear_tecnico(self, supervisor: dict, nombre: str, apellido: str, profesion: str = "mecanico") -> tuple[dict, dict]:
        """Lo registra su supervisor; devuelve el técnico y su cabecera de sesión"""
        respuesta = self.cliente.post("/api/v1/tecnicos", headers=supervisor, json={
            "nombre": nombre, "apellido_paterno": apellido, "profesion": profesion,
        })
        assert respuesta.status_code == 201, respuesta.text
        tecnico = respuesta.json()
        return tecnico, self.primer_acceso(tecnico["nombre_usuario"], "Tecn1co!")

    def crear_ubicacion(self) -> dict:
        respuesta = self.cliente.post("/api/v1/ubicaciones-tecnicas", headers=self.admin(), json={
            "sector": "Embotellado", "subsector": "Area 1", "sistema": "Llenadora",
        })
        assert respuesta.status_code == 201, respuesta.text
        return respuesta.json()


def plantilla_excel(personal, grupos, ubicaciones) -> bytes:
    """Excel con la forma de la plantilla de SAP: Personal en A3, Grupos en E3 y Ubicaciones en K7"""
    libro = Workbook()
    hoja = libro.active
    for j, titulo in enumerate(("Nombre Completo", "Grupo Planificador", "Supervisor"), start=1):
        hoja.cell(3, j, titulo)
    for i, fila in enumerate(personal, start=4):
        for j, valor in enumerate(fila, start=1):
            hoja.cell(i, j, valor)
    for j, titulo in enumerate(("Grupo", "Turno", "Supervisor"), start=5):
        hoja.cell(3, j, titulo)
    for i, fila in enumerate(grupos, start=4):
        for j, valor in enumerate(fila, start=5):
            hoja.cell(i, j, valor)
    for j, titulo in enumerate(("Sector", "Subsector", "Sistema", "Subsistema"), start=11):
        hoja.cell(7, j, titulo)
    for i, fila in enumerate(ubicaciones, start=8):
        for j, valor in enumerate(fila, start=11):
            hoja.cell(i, j, valor)
    archivo = BytesIO()
    libro.save(archivo)
    return archivo.getvalue()

from datetime import datetime, time
from io import BytesIO

import pytest
from openpyxl import load_workbook

from tests.integracion.conftest import plantilla_excel

pytestmark = pytest.mark.integracion

HOY = datetime.now().date()
INICIO = datetime.combine(HOY, time(8))
FIN = datetime.combine(HOY, time(16))
UBICACION = ("Embotellado", "Area 1", "Llenadora", "Sistema de Tapado")
ENCABEZADOS = ["OT", "Título", "Tipo", "Ubicación Técnica", "Técnico", "Prioridad", "Inicio", "Fin",
               "Equipo en Funcionamiento", "Clave", "Descripción del Paso", "Horas"]


def plantilla_con_ots(personal, ots) -> bytes:
    """La plantilla de SAP con Personal y Ubicaciones, más una hoja de órdenes de trabajo"""
    libro = load_workbook(BytesIO(plantilla_excel(personal, [], [UBICACION])))
    hoja = libro.create_sheet("Ordenes de trabajo")
    hoja.append(ENCABEZADOS)
    for fila in ots:
        hoja.append(fila)
    archivo = BytesIO()
    libro.save(archivo)
    return archivo.getvalue()


def ot(numero, titulo, tecnico, clave="PM01", descripcion="Cambiar tapadora", horas=2, ubicacion=" / ".join(UBICACION)):
    return [numero, titulo, "OE01", ubicacion, tecnico, 2, INICIO, FIN, "No", clave, descripcion, horas]


def paso(numero, clave, descripcion, horas=None):
    return [numero, None, None, None, None, None, None, None, None, clave, descripcion, horas]


@pytest.fixture
def supervisor(api):
    datos, sesion = api.crear_supervisor("Fernando", "Rojas", "Grupo 1")
    return {"sesion": sesion, "nombre": f"{datos['nombre']} {datos['apellido_paterno']}"}


def importar(api, sesion, contenido):
    respuesta = api.cliente.post(
        "/api/v1/importacion/mi-grupo", headers=sesion,
        files={"archivo": ("plantilla.xlsx", contenido, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")},
    )
    assert respuesta.status_code == 200, respuesta.text
    return respuesta.json()


def test_el_supervisor_importa_tecnicos_y_sus_ots_en_un_solo_archivo(api, supervisor):
    contenido = plantilla_con_ots(
        personal=[("Carlos Mamani Quispe", "Mecanico", supervisor["nombre"])],
        ots=[
            ot(1, "Cambio de tapadora", "Carlos Mamani Quispe", descripcion="Desmontar tapadora", horas=1.5),
            paso(1, "PM01", "Montar tapadora nueva", 2),
            paso(1, "PMNN", "Limpiar zona de trabajo"),
        ],
    )

    resultado = importar(api, supervisor["sesion"], contenido)

    assert resultado["errores"] == []
    assert len(resultado["usuarios_creados"]) == 1
    assert resultado["ots_creadas"] == 1
    ots = api.cliente.get("/api/v1/ordenes-trabajo/grupo", headers=supervisor["sesion"], params={"fecha": HOY.isoformat()}).json()
    assert [o["titulo"] for o in ots] == ["Cambio de tapadora"]
    assert [(p["numero_paso"], p["clave_control"]) for p in ots[0]["pasos"]] == [
        (None, "PMNN"), (None, "PMNN"), (None, "PMNN"), (1, "PM01"), (2, "PM01"), (3, "PMNN"),
    ]
    assert ots[0]["tecnico"]["nombre"] == "Carlos Mamani Quispe"


def test_cada_ot_con_problemas_se_reporta_y_las_demas_se_crean(api, supervisor):
    _, sesion_otro = api.crear_supervisor("Patricia", "Vargas", "Grupo 2")
    api.crear_tecnico(sesion_otro, "Erika", "Villca")
    contenido = plantilla_con_ots(
        personal=[("Carlos Mamani Quispe", "Mecanico", supervisor["nombre"])],
        ots=[
            ot(1, "Correcta", "Carlos Mamani Quispe"),
            ot(2, "De otro grupo", "Erika Villca"),
            ot(3, "Ubicación inexistente", "Carlos Mamani Quispe", ubicacion="Patio / Galpón 9"),
            ot(4, "Sin PM01", "Carlos Mamani Quispe", clave="PMNN", horas=None),
        ],
    )

    resultado = importar(api, supervisor["sesion"], contenido)

    assert resultado["ots_creadas"] == 1
    motivos = {e["dato"]: e["motivo"] for e in resultado["errores"]}
    assert motivos == {
        "OT 2: De otro grupo": "Erika Villca no es un técnico de tu grupo",
        "OT 3: Ubicación inexistente": "Ubicación técnica no registrada: Patio / Galpón 9",
        "OT 4: Sin PM01": "Se requiere al menos un paso PM01",
    }


def test_importar_dos_veces_las_mismas_ots_no_las_duplica(api, supervisor):
    contenido = plantilla_con_ots(
        personal=[("Carlos Mamani Quispe", "Mecanico", supervisor["nombre"])],
        ots=[ot(1, "Cambio de tapadora", "Carlos Mamani Quispe")],
    )
    importar(api, supervisor["sesion"], contenido)

    segunda = importar(api, supervisor["sesion"], contenido)

    assert segunda["ots_creadas"] == 0
    assert {"dato": "OT 1: Cambio de tapadora", "motivo": "Ya está registrada"} in [
        {"dato": o["dato"], "motivo": o["motivo"]} for o in segunda["omitidos"]
    ]

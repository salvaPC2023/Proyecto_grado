from datetime import datetime, time
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import text

pytestmark = pytest.mark.integracion

HOY = datetime.now().date()


@pytest.fixture
def grupo(api):
    """Un supervisor con un técnico en su grupo y una ubicación técnica, todos ya con su primer acceso"""
    _, supervisor = api.crear_supervisor("Fernando", "Rojas", "Grupo 1")
    tecnico, sesion_tecnico = api.crear_tecnico(supervisor, "Carlos", "Mamani")
    return {"supervisor": supervisor, "tecnico": tecnico, "sesion_tecnico": sesion_tecnico, "ubicacion": api.crear_ubicacion()}


def crear_ot(api, grupo, pasos=None) -> dict:
    respuesta = api.cliente.post("/api/v1/ordenes-trabajo", headers=grupo["supervisor"], json={
        "titulo": "Cambio de rodamientos",
        "tipo_de_orden": "OE01",
        "ubicacion_tecnica_id": grupo["ubicacion"]["id"],
        "tecnico_asignado_id": grupo["tecnico"]["id"],
        "prioridad": 1,
        "estatus_equipo": False,
        "fecha_inic_planif": datetime.combine(HOY, time(0, 0)).isoformat(),
        "fecha_fin_planif": datetime.combine(HOY, time(23, 0)).isoformat(),
        "pasos": pasos or [
            {"descripcion": "Desmontar motor", "clave_control": "PM01", "horas_planificadas": "1.5"},
            {"descripcion": "Montar rodamientos", "clave_control": "PM01", "horas_planificadas": "2"},
        ],
    })
    assert respuesta.status_code == 201, respuesta.text
    return respuesta.json()


def cerrar(api, sesion, ot, paso, descripcion="Se cambió el rodamiento", cierre_id=None):
    return api.cliente.post(
        f"/api/v1/ordenes-trabajo/{ot['id']}/pasos/{paso['id']}/cierre", headers=sesion,
        json={"resultado_trabajo": "ejecutado", "tiempo_real_trabajado": "1.5",
              "descripcion_trabajo_realizado": descripcion, "id": str(cierre_id or uuid4())},
    )


def test_la_ot_se_guarda_con_sus_pasos_y_los_datos_de_coordinacion(api, grupo):
    ot = crear_ot(api, grupo)

    detalle = api.cliente.get(f"/api/v1/ordenes-trabajo/{ot['id']}", headers=grupo["supervisor"]).json()
    assert [p["numero_paso"] for p in detalle["pasos"]] == [None, None, None, 1, 2]  # 3 de seguridad, sin número
    assert detalle["tecnico"] == {"nombre": "Carlos Mamani", "profesion": "mecanico"}
    assert detalle["ubicacion"]["sistema"] == "Llenadora"
    assert "descripcion" not in detalle


def test_el_tecnico_cierra_todos_los_pm01_y_la_ot_queda_cerrada(api, grupo):
    ot = crear_ot(api, grupo)
    tecnico = grupo["sesion_tecnico"]

    mis_ots = api.cliente.get("/api/v1/ordenes-trabajo/mis-ots", headers=tecnico, params={"fecha": HOY.isoformat()}).json()
    assert [o["id"] for o in mis_ots] == [ot["id"]]

    pm01 = [p for p in ot["pasos"] if p["clave_control"] == "PM01"]
    primero = cerrar(api, tecnico, ot, pm01[0])
    assert primero.status_code == 201 and primero.json()["estatus"] == "en_progreso"

    segundo = cerrar(api, tecnico, ot, pm01[1], descripcion="Se montaron rodamientos nuevos")
    assert segundo.status_code == 201
    assert segundo.json()["estatus"] == "cerrada"
    assert segundo.json()["fecha_cierre"] is not None


def test_los_cierres_quedan_guardados_en_postgresql(api, grupo, sesiones):
    ot = crear_ot(api, grupo)
    for paso in (p for p in ot["pasos"] if p["clave_control"] == "PM01"):
        cerrar(api, grupo["sesion_tecnico"], ot, paso)

    with sesiones() as sesion:
        cierres = sesion.execute(text("SELECT count(*) FROM cierres_paso_ot")).scalar_one()
        estatus = sesion.execute(text("SELECT estatus FROM ordenes_de_trabajo")).scalar_one()
    assert cierres == 2
    assert estatus == "cerrada"


def test_un_cierre_reenviado_desde_la_cola_sin_conexion_no_se_duplica(api, grupo, sesiones):
    # tarjeta 27: el celular reenvía el cierre si no recibió la respuesta; lleva siempre el mismo id
    ot = crear_ot(api, grupo)
    paso = next(p for p in ot["pasos"] if p["clave_control"] == "PM01")
    cierre_id = uuid4()

    primero = cerrar(api, grupo["sesion_tecnico"], ot, paso, cierre_id=cierre_id)
    reenvio = cerrar(api, grupo["sesion_tecnico"], ot, paso, cierre_id=cierre_id)

    assert primero.status_code == 201
    assert reenvio.status_code == 201  # no es un error: devuelve la OT como quedó
    with sesiones() as sesion:
        assert sesion.execute(text("SELECT count(*) FROM cierres_paso_ot")).scalar_one() == 1


def test_un_segundo_cierre_distinto_del_mismo_paso_se_rechaza(api, grupo):
    ot = crear_ot(api, grupo)
    paso = next(p for p in ot["pasos"] if p["clave_control"] == "PM01")
    cerrar(api, grupo["sesion_tecnico"], ot, paso)

    assert cerrar(api, grupo["sesion_tecnico"], ot, paso).status_code == 409


def test_el_excel_exportado_contiene_la_ot_y_su_cierre(api, grupo):
    ot = crear_ot(api, grupo)
    paso = next(p for p in ot["pasos"] if p["clave_control"] == "PM01")
    cerrar(api, grupo["sesion_tecnico"], ot, paso, descripcion="Se cambió el rodamiento del lado acople")

    respuesta = api.cliente.get(
        "/api/v1/ordenes-trabajo/exportar", headers=grupo["supervisor"],
        params={"desde": HOY.isoformat(), "hasta": HOY.isoformat()},
    )

    assert respuesta.status_code == 200
    libro = load_workbook(BytesIO(respuesta.content))
    valores = {c for hoja in libro for fila in hoja.iter_rows(values_only=True) for c in fila if c is not None}
    assert "Cambio de rodamientos" in valores
    assert "Se cambió el rodamiento del lado acople" in valores
    assert "Carlos Mamani" in valores


def test_una_ot_vencida_se_cierra_sola_y_el_tecnico_ya_no_puede_cerrarla(api, grupo, sesiones):
    from datetime import timedelta

    from src.modulos.ordenes_trabajo.aplicacion.cerrar_vencidas import cerrar_ots_vencidas
    from src.modulos.ordenes_trabajo.infraestructura.repositorio_ordenes_trabajo import RepositorioOrdenesTrabajoSQL

    hace_tres_dias = HOY - timedelta(days=3)
    respuesta = api.cliente.post("/api/v1/ordenes-trabajo", headers=grupo["supervisor"], json={
        "titulo": "Lubricación olvidada", "tipo_de_orden": "OE01", "ubicacion_tecnica_id": grupo["ubicacion"]["id"],
        "tecnico_asignado_id": grupo["tecnico"]["id"], "prioridad": 3, "estatus_equipo": True,
        "fecha_inic_planif": datetime.combine(hace_tres_dias, time(8)).isoformat(),
        "fecha_fin_planif": datetime.combine(hace_tres_dias, time(12)).isoformat(),
        "pasos": [{"descripcion": "Lubricar rodamientos", "clave_control": "PM01", "horas_planificadas": "1"}],
    })
    ot = respuesta.json()
    paso = next(p for p in ot["pasos"] if p["clave_control"] == "PM01")

    # aunque la tarea todavía no haya corrido, el técnico ya no puede cerrarla
    assert cerrar(api, grupo["sesion_tecnico"], ot, paso).status_code == 422

    with sesiones() as sesion:
        assert cerrar_ots_vencidas(datetime.now(), RepositorioOrdenesTrabajoSQL(sesion)) == 1

    detalle = api.cliente.get(f"/api/v1/ordenes-trabajo/{ot['id']}", headers=grupo["supervisor"]).json()
    assert detalle["estatus"] == "cerrada"
    cierre = next(p for p in detalle["pasos"] if p["clave_control"] == "PM01")["cierre"]
    assert cierre["resultado_trabajo"] == "no_ejecutado"
    assert cierre["tiempo_real_trabajado"] in ("0.00", "0", 0)


def test_no_se_planifica_un_paso_de_mas_de_12_horas(api, grupo):
    respuesta = api.cliente.post("/api/v1/ordenes-trabajo", headers=grupo["supervisor"], json={
        "titulo": "Absurda", "tipo_de_orden": "OE01", "ubicacion_tecnica_id": grupo["ubicacion"]["id"],
        "tecnico_asignado_id": grupo["tecnico"]["id"], "prioridad": 3, "estatus_equipo": True,
        "fecha_inic_planif": datetime.combine(HOY, time(8)).isoformat(),
        "fecha_fin_planif": datetime.combine(HOY, time(12)).isoformat(),
        "pasos": [{"descripcion": "Paso", "clave_control": "PM01", "horas_planificadas": "999"}],
    })

    assert respuesta.status_code == 422
    assert "12" in respuesta.json()["detail"]

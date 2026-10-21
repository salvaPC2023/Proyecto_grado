from datetime import time
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import Workbook

from src.modulos.importacion.aplicacion.importar import SupervisorSinGrupo, importar_carga_inicial, importar_para_mi_grupo
from src.modulos.importacion.dominio.puertos import ConsultaPersonal, PlantillaInvalida
from src.modulos.importacion.dominio.reglas import horario_de_turno, normalizar, profesion_de, separar_nombre
from src.modulos.importacion.infraestructura.lector_excel import leer_plantilla
from tests.modulos.acceso_roles.repositorios_falsos import RepositorioSupervisoresFalso, RepositorioTecnicosFalso, RepositorioUsuariosFalso
from tests.modulos.ordenes_trabajo.repositorios_falsos import RepositorioUbicacionesFalso

# Los archivos Excel se crean en memoria con openpyxl, con la misma forma que la plantilla real
# (prueba_importacion.xlsx): Personal en A3, Grupos en E3 y Ubicaciones Técnicas en K7.

PERSONAL = [
    ("Carlos Mamani Quispe", "Mecanico", "Fernando Rojas Aguilar"),
    ("Luis Torrez Vargas", "Electronico", "Fernando Rojas Aguilar"),
    ("Erika Villca Nina", "Electrico", "Patricia Vargas Mendoza"),
]
GRUPOS = [
    ("Grupo 1", "Turno 1", "Fernando Rojas Aguilar"),
    ("Grupo 2", "Turno 3", "Patricia Vargas Mendoza"),
]
UBICACIONES = [
    ("Embotellado", "Area 1", "Llenadora", "Sistema de Dosificacion"),
    ("Embotellado", "Area 1", "Llenadora", "Sistema de Tapado"),
]


def _plantilla(personal=PERSONAL, grupos=GRUPOS, ubicaciones=UBICACIONES, encabezados_personal=("Nombre Completo", "Grupo Planificador", "Supervisor")) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Tecnicos"
    if personal is not None:
        hoja.append([])
        hoja.append(["Personal"])
        hoja.append(list(encabezados_personal))  # fila 3
        for i, fila in enumerate(personal, start=4):
            for j, valor in enumerate(fila, start=1):
                hoja.cell(i, j, valor)
    if grupos is not None:
        for j, titulo in enumerate(("Grupo", "Turno", "Supervisor"), start=5):
            hoja.cell(3, j, titulo)
        for i, fila in enumerate(grupos, start=4):
            for j, valor in enumerate(fila, start=5):
                hoja.cell(i, j, valor)
    if ubicaciones is not None:
        for j, titulo in enumerate(("Sector", "Subsector", "Sistema", "Subsistema"), start=11):
            hoja.cell(7, j, titulo)
        for i, fila in enumerate(ubicaciones, start=8):
            for j, valor in enumerate(fila, start=11):
                hoja.cell(i, j, valor)
    archivo = BytesIO()
    libro.save(archivo)
    return archivo.getvalue()


class ConsultaPersonalFalsa(ConsultaPersonal):
    """Lee lo ya registrado en los repositorios falsos, como la consulta SQL lee la base"""

    def __init__(self, usuarios: RepositorioUsuariosFalso, supervisores: RepositorioSupervisoresFalso):
        self._usuarios, self._supervisores = usuarios, supervisores

    @staticmethod
    def _nombre(usuario):
        return normalizar(" ".join(p for p in (usuario.nombre, usuario.apellido_paterno, usuario.apellido_materno) if p))

    def nombres_completos(self):
        return {self._nombre(u) for u in self._usuarios.usuarios.values()}

    def supervisores_por_nombre(self):
        return {self._nombre(self._usuarios.obtener_por_id(s.usuario_id)): s.id for s in self._supervisores.listar()}


class RepositorioTecnicosConGrupos(RepositorioTecnicosFalso):
    """El grupo del supervisor es el que le asignó la importación"""

    def __init__(self, supervisores: RepositorioSupervisoresFalso):
        super().__init__()
        self._supervisores = supervisores

    def obtener_grupo_del_supervisor(self, supervisor_usuario_id):
        supervisor = self._supervisores.obtener_por_usuario_id(supervisor_usuario_id)
        return self._supervisores.obtener_grupo(supervisor.id).id if supervisor else None


class Sistema:
    """Repositorios en memoria compartidos entre importaciones, como una base de datos"""

    def __init__(self):
        self.usuarios = RepositorioUsuariosFalso()
        self.supervisores = RepositorioSupervisoresFalso()
        self.tecnicos = RepositorioTecnicosConGrupos(self.supervisores)
        self.ubicaciones = RepositorioUbicacionesFalso([])
        self.admin_id = uuid4()

    def carga_inicial(self, contenido: bytes):
        personal, ubicaciones, grupos = leer_plantilla(contenido)
        return importar_carga_inicial(
            personal, ubicaciones, self.admin_id, ConsultaPersonalFalsa(self.usuarios, self.supervisores),
            self.usuarios, self.supervisores, self.tecnicos, self.ubicaciones, grupos=grupos,
        )

    def carga_del_supervisor(self, contenido: bytes, supervisor_usuario_id):
        personal, ubicaciones, _ = leer_plantilla(contenido)
        return importar_para_mi_grupo(
            personal, ubicaciones, supervisor_usuario_id, ConsultaPersonalFalsa(self.usuarios, self.supervisores),
            self.usuarios, self.supervisores, self.tecnicos, self.ubicaciones,
        )

    def grupo_de(self, nombre_supervisor: str):
        supervisor_id = ConsultaPersonalFalsa(self.usuarios, self.supervisores).supervisores_por_nombre()[normalizar(nombre_supervisor)]
        return self.supervisores.obtener_grupo(supervisor_id)


# --- lectura de la plantilla ---

def test_lee_las_tres_tablas_con_su_numero_de_fila():
    personal, ubicaciones, grupos = leer_plantilla(_plantilla())

    assert [(p.fila, p.nombre_completo) for p in personal] == [(4, "Carlos Mamani Quispe"), (5, "Luis Torrez Vargas"), (6, "Erika Villca Nina")]
    assert ubicaciones[0].subsistema == "Sistema de Dosificacion"
    assert [(g.grupo, g.turno) for g in grupos] == [("Grupo 1", "Turno 1"), ("Grupo 2", "Turno 3")]


def test_la_tabla_de_grupos_es_opcional():
    _, _, grupos = leer_plantilla(_plantilla(grupos=None))

    assert grupos == []


def test_encabezados_incorrectos_rechazan_el_archivo():
    with pytest.raises(PlantillaInvalida, match="Nombre Completo"):
        leer_plantilla(_plantilla(encabezados_personal=("Nombre Completo", "Profesion", "Supervisor")))


def test_un_archivo_que_no_es_excel_se_rechaza():
    with pytest.raises(PlantillaInvalida, match="Excel"):
        leer_plantilla(b"esto no es un Excel")


# --- reglas ---

@pytest.mark.parametrize("nombre_completo, esperado", [
    ("Wilson Choquehuanca", ("Wilson", "Choquehuanca", None)),
    ("Carlos Mamani Quispe", ("Carlos", "Mamani", "Quispe")),
    ("Juan Carlos Pérez López", ("Juan Carlos", "Pérez", "López")),  # las dos últimas palabras son los apellidos
    ("Wilson", None),
])
def test_separar_el_nombre_completo(nombre_completo, esperado):
    assert separar_nombre(nombre_completo) == esperado


@pytest.mark.parametrize("grupo_planificador, profesion", [
    ("Mecanico", "mecanico"),
    ("Mecánico", "mecanico"),
    ("Electronico", "electromecanico"),  # electrónico se registra como electromecánico
    ("Soldador", None),
])
def test_grupo_planificador_a_profesion(grupo_planificador, profesion):
    assert profesion_de(grupo_planificador) == profesion


@pytest.mark.parametrize("turno, horario", [
    ("Turno 1", (time(23), time(7))),
    ("turno 2", (time(7), time(15))),
    ("3", (time(15), time(23))),
    ("Turno 9", None),
    ("", None),
])
def test_turno_a_horario(turno, horario):
    assert horario_de_turno(turno) == horario


# --- carga inicial del administrador ---

def test_crea_supervisores_con_el_grupo_y_turno_de_la_tabla_y_sus_tecnicos():
    sistema = Sistema()

    resultado = sistema.carga_inicial(_plantilla())

    assert [u.rol for u in resultado.usuarios_creados].count("supervisor") == 2
    assert [u.rol for u in resultado.usuarios_creados].count("tecnico") == 3
    assert resultado.ubicaciones_creadas == 2
    assert resultado.errores == []
    grupo = sistema.grupo_de("Fernando Rojas Aguilar")
    assert (grupo.nombre_de_grupo, grupo.horario_entrada) == ("Grupo 1", time(23))
    assert sum(t.grupo_id == grupo.id for t in sistema.tecnicos.tecnicos) == 2


def test_supervisor_que_no_esta_en_la_tabla_de_grupos_va_al_turno_2():
    sistema = Sistema()

    sistema.carga_inicial(_plantilla(grupos=None))

    grupo = sistema.grupo_de("Fernando Rojas Aguilar")
    assert (grupo.nombre_de_grupo, grupo.horario_entrada) == ("Grupo Fernando Rojas", time(7))


def test_importar_dos_veces_no_duplica_nada():
    sistema = Sistema()
    sistema.carga_inicial(_plantilla())

    segunda = sistema.carga_inicial(_plantilla())

    assert segunda.usuarios_creados == []
    assert segunda.ubicaciones_creadas == 0
    assert len(segunda.omitidos) == len(PERSONAL) + len(UBICACIONES)
    assert len(sistema.ubicaciones.listar_ubicaciones()) == 2


def test_las_filas_invalidas_se_informan_y_las_demas_se_importan():
    con_errores = PERSONAL + [
        ("Pedro Paz Rios", "Soldador", "Fernando Rojas Aguilar"),  # profesión desconocida
        ("Ana Lima Soto", "Mecanico", ""),  # sin supervisor
    ]
    sistema = Sistema()

    resultado = sistema.carga_inicial(_plantilla(personal=con_errores))

    assert [u.rol for u in resultado.usuarios_creados].count("tecnico") == 3
    assert [e.fila for e in resultado.errores] == [7, 8]


def test_sin_personal_ni_ubicaciones_rechaza_el_archivo():
    with pytest.raises(PlantillaInvalida):
        Sistema().carga_inicial(_plantilla(personal=None, ubicaciones=None))


# --- carga del supervisor para su grupo ---

def test_el_supervisor_solo_registra_a_los_tecnicos_de_su_grupo():
    sistema = Sistema()
    # el administrador ya cargó a los dos supervisores (tabla Grupos), a Carlos y las ubicaciones
    sistema.carga_inicial(_plantilla(personal=PERSONAL[:1]))
    patricia = next(s for s in sistema.supervisores.listar() if sistema.grupo_de("Patricia Vargas Mendoza").supervisor_id == s.id)

    resultado = sistema.carga_del_supervisor(_plantilla(), patricia.usuario_id)

    assert [u.nombre_completo for u in resultado.usuarios_creados] == ["Erika Villca Nina"]
    # Carlos ya estaba registrado, Luis es del grupo de Fernando y las ubicaciones ya existían
    assert len(resultado.omitidos) == 2 + len(UBICACIONES)


def test_supervisor_sin_grupo_no_puede_importar():
    sistema = Sistema()

    with pytest.raises(SupervisorSinGrupo):
        sistema.carga_del_supervisor(_plantilla(), uuid4())

from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import uuid4

import pytest
from openpyxl import load_workbook

from src.modulos.ordenes_trabajo.aplicacion.exportar_ots import RangoInvalido, exportar_ots
from src.modulos.ordenes_trabajo.dominio.modelos import CierrePaso, OrdenDeTrabajo, PasoOT
from src.modulos.ordenes_trabajo.infraestructura.excel_ots import generar_excel
from src.modulos.ubicaciones_tecnicas.dominio.modelos import UbicacionTecnica
from tests.modulos.ordenes_trabajo.repositorios_falsos import RepositorioOrdenesTrabajoFalso, RepositorioTecnicosFalso

SUPERVISOR, GRUPO, TECNICO = uuid4(), uuid4(), uuid4()
LLENADORA = UbicacionTecnica(id=uuid4(), sector="Embotellado", subsector="Area 1", sistema="Llenadora")
CALDERO = UbicacionTecnica(id=uuid4(), sector="Suministros", subsector="Vapor", sistema="Caldero")


def _ot(inicio: datetime, fin: datetime, ubicacion=LLENADORA, titulo="Cambiar rodamiento", pasos=None) -> OrdenDeTrabajo:
    ot_id = uuid4()
    return OrdenDeTrabajo(
        id=ot_id, titulo=titulo, tipo_de_orden="OE02", ubicacion_tecnica_id=ubicacion.id, tecnico_asignado_id=TECNICO,
        creado_por_id=uuid4(), descripcion="d", prioridad=2, estatus_equipo=True, estatus="asignada",
        fecha_inic_planif=inicio, fecha_fin_planif=fin,
        pasos=pasos if pasos is not None else [
            PasoOT(id=uuid4(), ot_id=ot_id, numero_paso=1, descripcion="Piense de manera inteligente", clave_control="PMNN"),
            PasoOT(id=uuid4(), ot_id=ot_id, numero_paso=4, descripcion="Desmontar", clave_control="PM01", horas_planificadas=Decimal("1.5")),
        ],
    )


def _exportar(ots, desde=None, hasta=None, ubicacion_tecnica_id=None):
    repositorio_ots = RepositorioOrdenesTrabajoFalso(grupo_de_tecnico={TECNICO: GRUPO})
    for ot in ots:
        repositorio_ots.crear(ot)
    repositorio_tecnicos = RepositorioTecnicosFalso(grupo_de_supervisor={SUPERVISOR: GRUPO})
    return exportar_ots(SUPERVISOR, repositorio_ots, repositorio_tecnicos, desde, hasta, ubicacion_tecnica_id)


def _filas(contenido: bytes) -> list[dict]:
    """Lee el Excel generado como lo haría el supervisor: busca la fila de encabezados y devuelve
    cada fila de datos como diccionario. No depende de la fila en que empiece la tabla."""
    hoja = load_workbook(BytesIO(contenido)).worksheets[0]
    filas = list(hoja.iter_rows(values_only=True))
    inicio = next(i for i, fila in enumerate(filas) if fila and fila[0] == "Título OT")
    encabezados = filas[inicio]
    return [dict(zip(encabezados, fila)) for fila in filas[inicio + 1:] if any(v is not None for v in fila)]


# --- filtros de la exportación ---

def test_fecha_inicial_posterior_a_la_final_lanza_error():
    with pytest.raises(RangoInvalido):
        _exportar([], desde=date(2026, 10, 10), hasta=date(2026, 10, 1))


def test_entra_la_ot_vigente_algun_dia_del_rango():
    # la OT dura del 4 al 6: aunque empiece antes del rango, estuvo vigente el 5
    de_varios_dias = _ot(datetime(2026, 10, 4, 8), datetime(2026, 10, 6, 17))
    anterior = _ot(datetime(2026, 9, 1, 8), datetime(2026, 9, 1, 12))
    posterior = _ot(datetime(2026, 11, 1, 8), datetime(2026, 11, 1, 12))

    ots = _exportar([de_varios_dias, anterior, posterior], desde=date(2026, 10, 5), hasta=date(2026, 10, 5))

    assert ots == [de_varios_dias]


def test_filtra_por_ubicacion_tecnica():
    en_la_llenadora = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12))
    en_el_caldero = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12), ubicacion=CALDERO)

    ots = _exportar([en_la_llenadora, en_el_caldero], ubicacion_tecnica_id=CALDERO.id)

    assert ots == [en_el_caldero]


def test_supervisor_sin_grupo_no_exporta_nada():
    repositorio_tecnicos = RepositorioTecnicosFalso(grupo_de_supervisor={})

    assert exportar_ots(SUPERVISOR, RepositorioOrdenesTrabajoFalso(), repositorio_tecnicos) == []


# --- contenido del archivo ---

def test_una_fila_por_paso_pm01_sin_los_pasos_de_seguridad():
    ot = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12))

    filas = _filas(generar_excel([ot], {LLENADORA.id: LLENADORA}, {TECNICO: "Carlos Mamani Quispe"}))

    assert len(filas) == 1
    assert filas[0]["N.º paso"] == 4
    assert filas[0]["Descripción del paso"] == "Desmontar"


def test_datos_de_la_ot_ubicacion_y_tecnico():
    ot = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12))

    fila = _filas(generar_excel([ot], {LLENADORA.id: LLENADORA}, {TECNICO: "Carlos Mamani Quispe"}))[0]

    assert fila["Título OT"] == "Cambiar rodamiento"
    assert fila["Ubicación técnica"] == "Embotellado / Area 1 / Llenadora"
    assert fila["Técnico"] == "Carlos Mamani Quispe"
    assert fila["Horas planificadas"] == 1.5
    assert fila["Inicio planificado"] == datetime(2026, 10, 5, 8)


def test_paso_sin_cierre_figura_como_pendiente():
    ot = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12))

    fila = _filas(generar_excel([ot], {LLENADORA.id: LLENADORA}, {}))[0]

    assert fila["Resultado"] == "Pendiente"
    assert fila["Tiempo real (h)"] is None


def test_paso_cerrado_lleva_los_datos_del_cierre():
    ot = _ot(datetime(2026, 10, 5, 8), datetime(2026, 10, 5, 12))
    paso = ot.pasos[1]
    paso.cierre = CierrePaso(
        id=uuid4(), paso_ot_id=paso.id, fecha_hora_notificacion=datetime(2026, 10, 5, 11, 30),
        tiempo_real_trabajado=Decimal("2"), trabajo_finalizado=True, sin_trabajo_realizado=False,
        resultado_trabajo="ejecutado", descripcion_trabajo_realizado="Equipo intervenido: Motor",
    )

    fila = _filas(generar_excel([ot], {LLENADORA.id: LLENADORA}, {}))[0]

    assert fila["Resultado"] == "Ejecutado"
    assert fila["Tiempo real (h)"] == 2
    assert fila["Trabajo finalizado"] == "Sí"
    assert fila["Descripción del trabajo realizado"] == "Equipo intervenido: Motor"


def test_sin_ots_genera_solo_los_encabezados():
    assert _filas(generar_excel([], {}, {})) == []

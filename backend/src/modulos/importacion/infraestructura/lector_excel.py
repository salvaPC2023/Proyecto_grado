from io import BytesIO
from zipfile import BadZipFile

from openpyxl import load_workbook

from ..dominio.modelos import FilaPersonal, FilaUbicacion
from ..dominio.puertos import PlantillaInvalida
from ..dominio.reglas import normalizar

# encabezados de cada tabla de la plantilla de SAP, en orden de izquierda a derecha
ENCABEZADOS_PERSONAL = ["nombre completo", "grupo planificador", "supervisor"]
ENCABEZADOS_UBICACIONES = ["sector", "subsector", "sistema", "subsistema"]


def _texto(valor) -> str:
    return "" if valor is None else str(valor).strip()


def _leer_tabla(libro, encabezados: list[str]) -> list[tuple[int, list[str]]]:
    """Busca la fila de encabezados en cualquier hoja y devuelve las filas de abajo hasta la primera vacía"""
    for hoja in libro.worksheets:
        for fila in hoja.iter_rows():
            for celda in fila:
                if normalizar(_texto(celda.value)) != encabezados[0]:
                    continue
                columna = celda.column
                vecinos = [normalizar(_texto(hoja.cell(celda.row, columna + i).value)) for i in range(len(encabezados))]
                if vecinos != encabezados:
                    raise PlantillaInvalida(
                        f"Los encabezados de la tabla en la hoja {hoja.title}, fila {celda.row}, deben ser: "
                        + ", ".join(e.title() for e in encabezados)
                    )
                datos = []
                numero = celda.row + 1
                while True:
                    valores = [_texto(hoja.cell(numero, columna + i).value) for i in range(len(encabezados))]
                    if not any(valores):
                        return datos
                    datos.append((numero, valores))
                    numero += 1
    return []


def leer_plantilla(contenido: bytes) -> tuple[list[FilaPersonal], list[FilaUbicacion]]:
    try:
        libro = load_workbook(BytesIO(contenido), read_only=False, data_only=True)
    except (BadZipFile, KeyError, ValueError, OSError):
        raise PlantillaInvalida("El archivo no es un Excel válido (.xlsx)")

    personal = [FilaPersonal(numero, *valores) for numero, valores in _leer_tabla(libro, ENCABEZADOS_PERSONAL)]
    ubicaciones = [FilaUbicacion(numero, *valores) for numero, valores in _leer_tabla(libro, ENCABEZADOS_UBICACIONES)]
    return personal, ubicaciones

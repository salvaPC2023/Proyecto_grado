from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from uuid import UUID

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from src.modulos.ubicaciones_tecnicas.dominio.modelos import UbicacionTecnica

from ..dominio.modelos import OrdenDeTrabajo

# el orden de las columnas no se cambia: el supervisor copia la tabla a un Excel con macros que carga SAP
# anchos ajustados para que las 18 columnas se lean impresas en una hoja A4 horizontal
COLUMNAS = [
    ("Título OT", 26), ("Tipo de orden", 9), ("Prioridad", 9), ("Estado OT", 11), ("Equipo en marcha", 10),
    ("Ubicación técnica", 28), ("Técnico", 18), ("Inicio planificado", 16), ("Fin planificado", 16), ("Fecha de cierre OT", 16),
    ("N.º paso", 7), ("Descripción del paso", 28), ("Horas planificadas", 11),
    ("Resultado", 13), ("Tiempo real (h)", 10), ("Fecha de notificación", 16), ("Trabajo finalizado", 10), ("Descripción del trabajo realizado", 42),
]
COLUMNAS_FECHA = (8, 9, 10, 16)
COLUMNAS_HORAS = (13, 15)
# el texto largo baja de línea dentro de su celda en vez de cortarse o taparse con la columna vecina
COLUMNAS_TEXTO_LARGO = (1, 6, 7, 12, 18)
COLUMNA_RESULTADO = 14

FORMATO_FECHA = "dd/mm/yyyy hh:mm"
FORMATO_HORAS = "0.00"
ESTADOS = {"asignada": "Asignada", "en_progreso": "En progreso", "cerrada": "Cerrada"}
RESULTADOS = {"ejecutado": "Ejecutado", "no_ejecutado": "No ejecutado"}

# colores de la app (ColoresApp.principal) y tonos suaves para que se impriman bien
AZUL = "131D8C"
FONDO_ALTERNO = "F1F5F9"
COLORES_RESULTADO = {
    "Pendiente": ("FEF3C7", "92400E"),
    "Ejecutado": ("D1FAE5", "065F46"),
    "No ejecutado": ("FEE2E2", "991B1B"),
}
LINEA_FINA = Side(style="thin", color="CBD5E1")
LINEA_SEPARADORA = Side(style="medium", color="64748B")


@dataclass
class EncabezadoReporte:
    """Datos del bloque de título; si no se envían, el reporte sale solo con la tabla"""
    grupo: str | None = None
    supervisor: str | None = None
    desde: date | None = None
    hasta: date | None = None
    ubicacion: str | None = None


def _ruta(ubicacion: UbicacionTecnica | None) -> str:
    if ubicacion is None:
        return ""
    return " / ".join(n for n in (ubicacion.sector, ubicacion.subsector, ubicacion.sistema, ubicacion.subsistema) if n)


def _rango(encabezado: EncabezadoReporte) -> str:
    if not encabezado.desde and not encabezado.hasta:
        return "Todas las fechas"
    desde = encabezado.desde.strftime("%d/%m/%Y") if encabezado.desde else "el inicio"
    hasta = encabezado.hasta.strftime("%d/%m/%Y") if encabezado.hasta else "hoy"
    return f"Del {desde} al {hasta}"


def _escribir_titulo(hoja: Worksheet, encabezado: EncabezadoReporte) -> int:
    """Escribe el bloque de título sin combinar celdas y devuelve la fila donde empieza la tabla"""
    hoja.cell(1, 1, "Reporte de Órdenes de Trabajo — ESPODI").font = Font(bold=True, size=16, color=AZUL)
    datos = [
        ("Grupo", encabezado.grupo or "—"),
        ("Supervisor", encabezado.supervisor or "—"),
        ("Rango de fechas", _rango(encabezado)),
        ("Ubicación técnica", encabezado.ubicacion or "Todas"),
        ("Generado", datetime.now().strftime("%d/%m/%Y %H:%M")),
    ]
    for i, (etiqueta, valor) in enumerate(datos, start=2):
        hoja.cell(i, 1, etiqueta).font = Font(bold=True, color="475569")
        hoja.cell(i, 2, valor)
    return len(datos) + 3  # una fila en blanco entre el título y la tabla


def _filas_de_datos(ots: list[OrdenDeTrabajo], ubicaciones: dict[UUID, UbicacionTecnica], tecnicos: dict[UUID, str]):
    """Una fila por paso PM01 con todos los datos de su OT repetidos, para que cada fila se pueda copiar sola"""
    for numero_ot, ot in enumerate(ots):
        for paso in (p for p in ot.pasos if p.clave_control == "PM01"):
            cierre = paso.cierre
            yield numero_ot, [
                ot.titulo, ot.tipo_de_orden, ot.prioridad, ESTADOS.get(ot.estatus, ot.estatus), "Sí" if ot.estatus_equipo else "No",
                _ruta(ubicaciones.get(ot.ubicacion_tecnica_id)), tecnicos.get(ot.tecnico_asignado_id, ""),
                ot.fecha_inic_planif, ot.fecha_fin_planif, ot.fecha_cierre,
                paso.numero_paso, paso.descripcion, float(paso.horas_planificadas or 0),
                RESULTADOS.get(cierre.resultado_trabajo) if cierre else "Pendiente",
                float(cierre.tiempo_real_trabajado) if cierre else None,
                cierre.fecha_hora_notificacion if cierre else None,
                ("Sí" if cierre.trabajo_finalizado else "No") if cierre else None,
                cierre.descripcion_trabajo_realizado if cierre else None,
            ]


def _hoja_ordenes(hoja: Worksheet, ots, ubicaciones, tecnicos, encabezado: EncabezadoReporte | None) -> None:
    fila_encabezados = _escribir_titulo(hoja, encabezado) if encabezado else 1

    for columna, (nombre, ancho) in enumerate(COLUMNAS, start=1):
        celda = hoja.cell(fila_encabezados, columna, nombre)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor=AZUL)
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        celda.border = Border(left=LINEA_FINA, right=LINEA_FINA, top=LINEA_FINA, bottom=LINEA_FINA)
        hoja.column_dimensions[celda.column_letter].width = ancho
    hoja.row_dimensions[fila_encabezados].height = 32

    fila = fila_encabezados
    ot_anterior = None
    for numero_ot, valores in _filas_de_datos(ots, ubicaciones, tecnicos):
        fila += 1
        empieza_ot = numero_ot != ot_anterior and ot_anterior is not None
        ot_anterior = numero_ot
        # las OT se alternan de color; todos los pasos de una misma OT comparten el mismo fondo
        fondo = PatternFill("solid", fgColor=FONDO_ALTERNO) if numero_ot % 2 else None
        for columna, valor in enumerate(valores, start=1):
            celda = hoja.cell(fila, columna, valor)
            celda.border = Border(
                left=LINEA_FINA, right=LINEA_FINA, bottom=LINEA_FINA,
                top=LINEA_SEPARADORA if empieza_ot else LINEA_FINA,  # línea más gruesa donde empieza otra OT
            )
            celda.alignment = Alignment(vertical="top", wrap_text=columna in COLUMNAS_TEXTO_LARGO)
            if fondo:
                celda.fill = fondo
            if columna in COLUMNAS_FECHA:
                celda.number_format = FORMATO_FECHA
            if columna in COLUMNAS_HORAS:
                celda.number_format = FORMATO_HORAS
        relleno, letra = COLORES_RESULTADO[valores[COLUMNA_RESULTADO - 1]]
        resultado = hoja.cell(fila, COLUMNA_RESULTADO)
        resultado.fill = PatternFill("solid", fgColor=relleno)
        resultado.font = Font(bold=True, color=letra)

    ultima_columna = hoja.cell(fila_encabezados, len(COLUMNAS)).column_letter
    hoja.auto_filter.ref = f"A{fila_encabezados}:{ultima_columna}{max(fila, fila_encabezados)}"
    hoja.freeze_panes = f"A{fila_encabezados + 1}"  # el encabezado queda visible al bajar
    _configurar_impresion(hoja, fila_encabezados)


def _hoja_resumen(hoja: Worksheet, ots: list[OrdenDeTrabajo], tecnicos: dict[UUID, str]) -> None:
    negrita = Font(bold=True)
    hoja.cell(1, 1, "Resumen").font = Font(bold=True, size=16, color=AZUL)

    hoja.cell(3, 1, "OT por estado").font = negrita
    for columna, titulo in enumerate(("Estado", "Cantidad"), start=1):
        _celda_encabezado(hoja.cell(4, columna, titulo))
    fila = 4
    for clave, nombre in ESTADOS.items():
        fila += 1
        hoja.cell(fila, 1, nombre)
        hoja.cell(fila, 2, sum(1 for ot in ots if ot.estatus == clave))
    fila += 1
    hoja.cell(fila, 1, "Total").font = negrita
    hoja.cell(fila, 2, len(ots)).font = negrita

    fila += 2
    hoja.cell(fila, 1, "Horas por técnico (pasos PM01)").font = negrita
    fila += 1
    for columna, titulo in enumerate(("Técnico", "Horas planificadas", "Horas reales", "Diferencia (real − plan.)", "Pasos cerrados"), start=1):
        _celda_encabezado(hoja.cell(fila, columna, titulo))
    # técnico -> [planificadas, reales, pasos cerrados, pasos totales]
    por_tecnico: dict[UUID, list] = {}
    for ot in ots:
        totales = por_tecnico.setdefault(ot.tecnico_asignado_id, [Decimal(0), Decimal(0), 0, 0])
        for paso in (p for p in ot.pasos if p.clave_control == "PM01"):
            totales[0] += paso.horas_planificadas or 0
            totales[3] += 1
            if paso.cierre:
                totales[1] += paso.cierre.tiempo_real_trabajado
                totales[2] += 1
    for tecnico_id, (planificadas, reales, cerrados, total) in sorted(por_tecnico.items(), key=lambda t: tecnicos.get(t[0], "")):
        fila += 1
        hoja.cell(fila, 1, tecnicos.get(tecnico_id, "—"))
        for columna, valor in ((2, planificadas), (3, reales), (4, reales - planificadas)):
            hoja.cell(fila, columna, float(valor)).number_format = FORMATO_HORAS
        hoja.cell(fila, 5, f"{cerrados} de {total}")

    for letra, ancho in (("A", 32), ("B", 18), ("C", 14), ("D", 24), ("E", 15)):
        hoja.column_dimensions[letra].width = ancho
    _configurar_impresion(hoja, None)


def _celda_encabezado(celda) -> None:
    celda.font = Font(bold=True, color="FFFFFF")
    celda.fill = PatternFill("solid", fgColor=AZUL)


def _configurar_impresion(hoja: Worksheet, fila_encabezados: int | None) -> None:
    """Horizontal y ajustada al ancho de la hoja; el encabezado de la tabla se repite en cada página"""
    hoja.page_setup.orientation = "landscape"
    hoja.page_setup.paperSize = hoja.PAPERSIZE_A4
    hoja.page_setup.fitToWidth = 1
    hoja.page_setup.fitToHeight = 0  # 0: las páginas que hagan falta a lo alto
    hoja.sheet_properties.pageSetUpPr.fitToPage = True
    hoja.print_options.horizontalCentered = True
    # márgenes estrechos (en pulgadas) para aprovechar el ancho de la hoja
    hoja.page_margins.left = hoja.page_margins.right = 0.25
    hoja.page_margins.top = hoja.page_margins.bottom = 0.5
    if fila_encabezados:
        hoja.print_title_rows = f"{fila_encabezados}:{fila_encabezados}"


def generar_excel(
    ots: list[OrdenDeTrabajo],
    ubicaciones: dict[UUID, UbicacionTecnica],
    tecnicos: dict[UUID, str],
    encabezado: EncabezadoReporte | None = None,
) -> bytes:
    """Hoja "Ordenes de trabajo": una fila por cada paso PM01 (los pasos de seguridad PMNN no se exportan).
    Hoja "Resumen": OT por estado y horas planificadas contra reales por técnico"""
    libro = Workbook()
    _hoja_ordenes(libro.active, ots, ubicaciones, tecnicos, encabezado)
    libro.active.title = "Ordenes de trabajo"
    _hoja_resumen(libro.create_sheet("Resumen"), ots, tecnicos)

    archivo = BytesIO()
    libro.save(archivo)
    return archivo.getvalue()
